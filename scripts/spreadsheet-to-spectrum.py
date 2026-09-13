import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WHEELS = os.path.join(ROOT, "vendor", "wheels")
for wheel in ("xlrd-2.0.2-py2.py3-none-any.whl", "et_xmlfile-2.0.0-py3-none-any.whl", "openpyxl-3.1.5-py2.py3-none-any.whl"):
    sys.path.insert(0, os.path.join(WHEELS, wheel))


def number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(str(value).strip())
        return result if result == result and abs(result) != float("inf") else None
    except (TypeError, ValueError):
        return None


def format_number(value):
    return format(value, ".15g")


def extract(sheets):
    best, best_expected = [], None
    for rows in sheets:
        found, expected = [], None
        for row in rows:
            a = row[0] if len(row) > 0 else None
            b = row[1] if len(row) > 1 else None
            if isinstance(a, str) and a.strip().upper().startswith("SPECTRTXT="):
                try:
                    expected = int(a.split("=", 1)[1].strip())
                except ValueError:
                    pass
            channel, counts = number(a), number(b)
            if channel is not None and counts is not None:
                found.append((channel, counts))
        if len(found) > len(best):
            best, best_expected = found, expected
    if len(best) < 3:
        raise ValueError("未找到至少三行的通道—计数两列数值表。")
    if best_expected is not None and best_expected != len(best):
        print(f"EXPECTED_COUNT={best_expected};ACTUAL_COUNT={len(best)}", file=sys.stderr)
    return "\n".join(f"{format_number(a)},{format_number(b)}" for a, b in best)


def read_xls(path):
    import xlrd
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        return [[sheet.row_values(i) for i in range(sheet.nrows)] for sheet in book.sheets()]
    finally:
        book.release_resources()


def read_xlsx(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(sheet.iter_rows(values_only=True)) for sheet in book.worksheets]
    finally:
        book.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("用法：spreadsheet-to-spectrum.py <xls|xlsx>")
    source = os.path.abspath(sys.argv[1])
    extension = os.path.splitext(source)[1].lower()
    try:
        sheets = read_xls(source) if extension == ".xls" else read_xlsx(source)
        sys.stdout.write(extract(sheets))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(2)
